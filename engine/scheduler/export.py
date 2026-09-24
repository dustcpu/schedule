# -*- coding: utf-8 -*-
"""导出层：写 result.xlsx / result.pdf / status.json（严格遵守接口协议 v1.1）。

result.xlsx：第一个 sheet 固定为「总览」，之后每个方案一个 sheet「方案N」。
result.pdf：第一页总览，之后每个方案一页（班级多则自动分页）。
status.json：UTF-8 无 BOM，code / message / plans / warnings。
"""
import os
import json
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from .config import DAYS, NUM_DAYS
from .data.load import Problem
from .solver.solve import Plan

HEADER_FILL = "2563EB"
SUBHEADER_FILL = "E5E7EB"
CLASSES_PER_PAGE = 3  # PDF 每页放几个班的课表


# ---------------------------------------------------------------- 字体
def register_font() -> str:
    """注册中文字体，失败退回 Helvetica（PDF 不允许中文乱码，见协议 §4）。"""
    for name, path in (
        ("MSYH", "C:/Windows/Fonts/msyh.ttc"),
        ("SimSun", "C:/Windows/Fonts/simsun.ttc"),
        ("SimHei", "C:/Windows/Fonts/simhei.ttf"),
    ):
        try:
            if os.path.exists(path):
                pdfmetrics.registerFont(TTFont(name, path))
                return name
        except Exception:
            continue
    return "Helvetica"


# ---------------------------------------------------------------- 课表网格
def _grid_rows(plan: Plan, ci, cfg) -> List[List[str]]:
    """生成某班课表的二维表（含表头）。"""
    labels = cfg.period_labels()
    rows: List[List[str]] = [["时间"] + DAYS]
    for per in cfg.periods():
        row = [labels[per - 1]]
        for d in range(1, NUM_DAYS + 1):
            row.append(plan.grid.get((ci.id, d, per), ""))
        rows.append(row)
    return rows


# ---------------------------------------------------------------- Excel
def write_xlsx(path: str, plans: List[Plan], problem: Problem) -> None:
    cfg = problem.config
    wb = Workbook()

    # 总览 sheet（协议：第一个 sheet 固定为「总览」）
    ws = wb.active
    ws.title = "总览"
    thin = Side(style="thin", color="9CA3AF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["编号", "方案名", "评分", "特点"]
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
    for i, p in enumerate(plans, start=2):
        vals = [p.index, p.name, p.score, p.note]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.alignment = Alignment(horizontal="center" if j != 4 else "left",
                                    vertical="center")
            c.border = border
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 60

    # 每个方案一个 sheet
    for p in plans:
        w = wb.create_sheet(f"方案{p.index}")
        r = 1
        for ci in problem.classes:
            title = w.cell(row=r, column=1, value=f"{ci.name or ci.id}　（{ci.track or '—'}）")
            title.font = Font(bold=True, size=12)
            r += 1
            rows = _grid_rows(p, ci, cfg)
            for ri, row in enumerate(rows):
                for cj, val in enumerate(row, start=1):
                    c = w.cell(row=r + ri, column=cj, value=val)
                    c.alignment = Alignment(horizontal="center", vertical="center")
                    c.border = border
                    if ri == 0:
                        c.font = Font(bold=True)
                        c.fill = PatternFill("solid", fgColor=SUBHEADER_FILL)
            r += len(rows) + 2  # 班级之间空两行
        w.column_dimensions["A"].width = 18
        for col in "BCDEF":
            w.column_dimensions[col].width = 14

    wb.save(path)


# ---------------------------------------------------------------- PDF
def write_pdf(path: str, task_id: str, plans: List[Plan], problem: Problem,
              requirements: str) -> None:
    cfg = problem.config
    font = register_font()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("T", parent=styles["Title"], fontName=font,
                                 fontSize=16, leading=22)
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontName=font,
                        fontSize=12, leading=16)
    body = ParagraphStyle("B", parent=styles["Normal"], fontName=font,
                          fontSize=9, leading=13)

    doc = SimpleDocTemplate(path, pagesize=A4,
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm)
    story: List[Any] = []

    # 第 1 页：总览
    story.append(Paragraph("排课结果总览", title_style))
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph(f"任务 ID：{task_id}", body))
    story.append(Paragraph(f"班级数：{len(problem.classes)}　　方案数：{len(plans)}", body))
    if requirements:
        story.append(Spacer(1, 3 * mm))
        story.append(Paragraph(f"特殊要求：{requirements[:60]}", body))
    story.append(Spacer(1, 5 * mm))

    data = [["编号", "方案名", "评分", "特点"]]
    for p in plans:
        data.append([str(p.index), p.name, f"{p.score:.1f}", p.note])
    t = Table(data, colWidths=[16 * mm, 32 * mm, 16 * mm, 104 * mm])
    t.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), font),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + HEADER_FILL)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 0), (2, -1), "CENTER"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)

    # 每方案若干页
    for p in plans:
        chunks = [problem.classes[i:i + CLASSES_PER_PAGE]
                  for i in range(0, len(problem.classes), CLASSES_PER_PAGE)] or [[]]
        for ci, group in enumerate(chunks):
            story.append(PageBreak())
            suffix = f"（{ci + 1}/{len(chunks)}）" if len(chunks) > 1 else ""
            story.append(Paragraph(f"方案{p.index}：{p.name}　评分 {p.score:.1f}{suffix}",
                                   title_style))
            story.append(Spacer(1, 4 * mm))
            for c in group:
                story.append(Paragraph(f"{c.name or c.id}　（{c.track or '—'}）", h2))
                rows = _grid_rows(p, c, cfg)
                tbl = Table(rows, colWidths=[34 * mm] + [29 * mm] * NUM_DAYS)
                tbl.setStyle(TableStyle([
                    ("FONTNAME", (0, 0), (-1, -1), font),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + SUBHEADER_FILL)),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]))
                story.append(tbl)
                story.append(Spacer(1, 5 * mm))

    doc.build(story)


# ---------------------------------------------------------------- status.json
def write_status(path: str, task_id: str, code: int, message: str,
                 plans: List[Plan], warnings: List[str]) -> None:
    payload = {
        "task_id": task_id,
        "code": code,
        "message": message,
        "plans": [
            {"index": p.index, "name": p.name, "score": p.score, "note": p.note}
            for p in plans
        ],
        "warnings": warnings,
    }
    # UTF-8 无 BOM（协议 §7）
    with open(path, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
