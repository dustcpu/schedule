# -*- coding: utf-8 -*-
"""数据加载层：读 input.xlsx、hard_limits.json、requirements.txt。

支持两种输入格式（自动检测）：

【新格式 · 推荐】4个 sheet：
  Sheet「教师」：教师ID | 教师姓名 | 任教学科 | 职务
  Sheet「班级」：班级ID | 班级名称 | 年级 | 班主任 | 选科（如物化生）
  Sheet「课时标准」：学科 | 选考周课时 | 非选考周课时 | 连堂节数 | 连堂日
  Sheet「任课安排」：教师ID | 班级ID | 学科（可选，默认用教师任教学科）

【旧格式 · 兼容】3个 sheet：
  Sheet「班级」：班级ID | 班级名称 | 年级 | 科类
  Sheet「教师」：教师ID | 教师姓名 | 任教学科
  Sheet「课程」：班级ID | 学科 | 教师ID | 周课时 | 连堂节数 | 单双周
"""
import json
import os
import re

# 这些科目在课表上不显示教师姓名，也不需要在任课安排中指定
NO_TEACHER_SUBJECTS = {"体育", "体育活动", "艺术", "音乐", "美术", "信息技术", "通用技术", "心理", "班会", "研究性学习", "校本课", "自习"}
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Any

from openpyxl import load_workbook

from ..config import (
    Config, SELF_STUDY, _norm_header,
    load_from_hard_limits,
)


class DataError(Exception):
    """输入数据错误，交由上层转成 code=2。"""


@dataclass
class ClassInfo:
    id: str
    name: str = ""
    grade: str = ""
    track: str = ""
    elective: str = ""


@dataclass
class TeacherInfo:
    id: str
    name: str = ""
    subject: str = ""
    duty: str = ""


@dataclass
class PeriodStandard:
    subject: str
    elective_periods: int = 0
    non_elective_periods: int = 0
    block_len: int = 0
    consec_day: int = 0


@dataclass
class CourseReq:
    class_id: str
    subject: str
    teacher_id: Optional[str]
    weekly: int
    block_len: int = 0
    week_mode: str = "每周"


@dataclass
class Problem:
    classes: List[ClassInfo] = field(default_factory=list)
    teachers: Dict[str, TeacherInfo] = field(default_factory=dict)
    courses: List[CourseReq] = field(default_factory=list)
    requirements: str = ""
    config: Config = field(default_factory=Config)
    warnings: List[str] = field(default_factory=list)

    def teacher_name(self, tid: Optional[str]) -> str:
        if not tid:
            return ""
        t = self.teachers.get(tid)
        return t.name if t else tid


def _find_sheet(wb, names):
    norm = {_norm_header(s): s for s in wb.sheetnames}
    for n in names:
        if _norm_header(n) in norm:
            return wb[norm[_norm_header(n)]]
    return None


def _read_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_norm_header(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        d = {}
        for i, h in enumerate(headers):
            if not h:
                continue
            d[h] = r[i] if i < len(r) else None
        out.append(d)
    return out


def _cell_str(v):
    return "" if v is None else str(v).strip()


def _cell_int(v, default=0):
    if v is None or str(v).strip() == "":
        return default
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return default


def _load_classes(ws):
    if ws is None:
        raise DataError("input.xlsx 缺少「班级」sheet")
    rows = _read_rows(ws)
    if not rows:
        raise DataError("「班级」sheet 为空，至少需要 1 个班级")
    out = []
    seen = set()
    for r in rows:
        cid = _cell_str(r.get("班级ID") or r.get("班级编号"))
        if not cid:
            continue
        if cid in seen:
            raise DataError(f"「班级」sheet 中班级ID 重复：{cid}")
        seen.add(cid)
        out.append(ClassInfo(
            id=cid,
            name=_cell_str(r.get("班级名称") or r.get("班级")),
            grade=_cell_str(r.get("年级")),
            track=_cell_str(r.get("科类")),
            elective=_cell_str(r.get("选科")),
        ))
    if not out:
        raise DataError("「班级」sheet 未解析到有效行（缺少「班级ID」列？）")
    return out


def _load_teachers(ws):
    if ws is None:
        return {}
    out = {}
    for r in _read_rows(ws):
        tid = _cell_str(r.get("教师ID") or r.get("教师编号"))
        if not tid:
            continue
        out[tid] = TeacherInfo(
            id=tid,
            name=_cell_str(r.get("教师姓名") or r.get("姓名")) or tid,
            subject=_cell_str(r.get("任教学科") or r.get("学科")),
            duty=_cell_str(r.get("职务")),
        )
    return out


def _load_period_standards(ws):
    if ws is None:
        return []
    out = []
    for r in _read_rows(ws):
        subj = _cell_str(r.get("学科") or r.get("科目"))
        if not subj:
            continue
        out.append(PeriodStandard(
            subject=subj,
            elective_periods=_cell_int(r.get("选考周课时") or r.get("选考课时"), 0),
            non_elective_periods=_cell_int(r.get("非选考周课时") or r.get("非选考课时"), 0),
            block_len=_cell_int(r.get("连堂节数"), 0),
            consec_day=_cell_int(r.get("连堂日"), 0),
        ))
    return out


def _load_teaching_assignments(ws):
    if ws is None:
        return []
    out = []
    for r in _read_rows(ws):
        tid = _cell_str(r.get("教师ID") or r.get("教师编号"))
        cid = _cell_str(r.get("班级ID") or r.get("班级编号"))
        if not tid or not cid:
            continue
        out.append({
            "teacher_id": tid,
            "class_id": cid,
            "subject": _cell_str(r.get("学科") or r.get("科目")),
        })
    return out


def _load_courses_old(ws, class_ids):
    if ws is None:
        raise DataError("input.xlsx 缺少「课程」sheet")
    rows = _read_rows(ws)
    if not rows:
        raise DataError("「课程」sheet 为空，无法排课")
    known = set(class_ids)
    out = []
    seen = set()
    for r in rows:
        cid = _cell_str(r.get("班级ID") or r.get("班级编号"))
        subj = _cell_str(r.get("学科") or r.get("科目"))
        if not cid or not subj:
            continue
        if cid not in known:
            raise DataError(f"「课程」sheet 出现未知班级ID：{cid}")
        key = (cid, subj)
        if key in seen:
            raise DataError(f"「课程」sheet 中 (班级 {cid}, 学科 {subj}) 重复")
        seen.add(key)
        tid = _cell_str(r.get("教师ID") or r.get("教师编号"))
        weekly = _cell_int(r.get("周课时"), 0)
        if weekly <= 0:
            raise DataError(f"班级 {cid} 的「{subj}」周课时应为正整数")
        block_len = _cell_int(r.get("连堂节数"), 0)
        wm = _cell_str(r.get("单双周")) or "每周"
        out.append(CourseReq(
            class_id=cid, subject=subj, teacher_id=tid or None,
            weekly=weekly, block_len=block_len if block_len > 0 else 0, week_mode=wm,
        ))
    if not out:
        raise DataError("「课程」sheet 未解析到有效行")
    return out


def _parse_teacher_constraints(text, teachers, classes):
    """从额外约束文本中解析教师指定，返回 {(班级ID, 学科): 教师ID}。
    支持：T001=C01语文 / T001教C01语文 / 张老师教高二1班语文
    """
    result = {}
    if not text:
        return result
    text = text.lstrip('\ufeff')
    name_to_tid = {t.name: tid for tid, t in teachers.items()}
    cname_to_cid = {c.name: c.id for c in classes}
    all_subjects = set(t.subject for t in teachers.values() if t.subject)
    class_ids = {c.id for c in classes}
    for line in text.splitlines():
        line = line.strip().lstrip('\ufeff')
        if not line:
            continue
        line = re.sub(r'^[指定：\s]+', '', line)
        m = re.match(r'^(.+?)\s*[=教]\s*(.+)$', line)
        if not m:
            continue
        t_str = m.group(1).strip()
        rest = m.group(2).strip()
        tid = t_str if t_str in teachers else name_to_tid.get(t_str)
        if not tid:
            continue
        cid = None
        subj = rest
        cm = re.match(r'^(C\d+)\s*(.*)$', rest)
        if cm:
            cid = cm.group(1)
            subj = cm.group(2).strip()
        else:
            for cname, cid_tmp in cname_to_cid.items():
                if rest.startswith(cname):
                    cid = cid_tmp
                    subj = rest[len(cname):].strip()
                    break
        if not cid or cid not in class_ids:
            continue
        matched = None
        for s in sorted(all_subjects, key=len, reverse=True):
            if s in subj:
                matched = s
                break
        if matched:
            result[(cid, matched)] = tid
    return result


def _generate_courses_new(classes, teachers, standards, assignments, warnings, teacher_constraints=None):
    std_map = {s.subject: s for s in standards}
    assignment_map = {}
    for a in assignments:
        tid = a["teacher_id"]
        cid = a["class_id"]
        subj = a["subject"]
        if not subj and tid in teachers:
            subj = teachers[tid].subject
        if subj:
            assignment_map[(cid, subj)] = tid
    if teacher_constraints:
        assignment_map.update(teacher_constraints)

    teachers_by_subject = {}
    for tid, t in teachers.items():
        if t.subject:
            teachers_by_subject.setdefault(t.subject, []).append(tid)

    teacher_load = {tid: 0 for tid in teachers}
    courses = []
    auto_assigned = 0
    subj_map = {"物": "物理", "化": "化学", "生": "生物", "政": "政治", "史": "历史", "地": "地理"}

    for ci in classes:
        elective_set = set()
        if ci.elective:
            for ch in ci.elective:
                if ch in subj_map:
                    elective_set.add(subj_map[ch])
        for std in standards:
            subj = std.subject
            is_elective = subj in elective_set
            weekly = std.elective_periods if is_elective else std.non_elective_periods
            if weekly <= 0:
                continue
            if subj in NO_TEACHER_SUBJECTS:
                tid = None
            else:
                tid = assignment_map.get((ci.id, subj))
                if not tid:
                    candidates = teachers_by_subject.get(subj, [])
                    if candidates:
                        tid = min(candidates, key=lambda t: teacher_load.get(t, 0))
                        teacher_load[tid] = teacher_load.get(tid, 0) + weekly
                        auto_assigned += 1
                    else:
                        tid = None
            courses.append(CourseReq(
                class_id=ci.id, subject=subj, teacher_id=tid,
                weekly=weekly, block_len=std.block_len, week_mode="每周",
            ))
    if auto_assigned > 0:
        warnings.append(
            f"有 {auto_assigned} 条课程未指定教师，已按学科和工作量自动分配。"
            f"如需特殊指定，可在「特殊要求」中输入（格式：T001教C01语文）。"
        )
    return courses



def load_problem(in_dir):
    xlsx_path = None
    for fn in sorted(os.listdir(in_dir)):
        if fn.lower().startswith("input.") and fn.lower().endswith((".xlsx", ".xlsm")):
            xlsx_path = os.path.join(in_dir, fn)
            break
    if xlsx_path is None:
        raise DataError("输入目录里没有 input.xlsx")

    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    classes = _load_classes(_find_sheet(wb, ["班级", "班级表", "classes"]))
    teachers = _load_teachers(_find_sheet(wb, ["教师", "教师表", "teachers"]))

    period_ws = _find_sheet(wb, ["课时标准", "课时", "period_standards"])
    is_new_format = period_ws is not None
    warnings_list = []

    if is_new_format:
        standards = _load_period_standards(period_ws)
        assignments = _load_teaching_assignments(
            _find_sheet(wb, ["任课安排", "任课", "teaching_assignments"]))
        if not standards:
            raise DataError("「课时标准」sheet 为空，无法排课")
        courses = _generate_courses_new(classes, teachers, standards, assignments, warnings_list)
        warnings_list.append(f"新格式：{len(standards)} 门学科标准，生成 {len(courses)} 条课程")
    else:
        courses = _load_courses_old(
            _find_sheet(wb, ["课程", "课程表", "课时", "courses"]),
            [c.id for c in classes])

    wb.close()

    cfg = Config()
    hl_path = os.path.join(in_dir, "hard_limits.json")
    if os.path.exists(hl_path):
        try:
            with open(hl_path, "r", encoding="utf-8-sig") as f:
                hl = json.load(f)
            cfg = load_from_hard_limits(hl)
        except Exception as e:
            cfg = Config()
            warnings_list.append(f"hard_limits.json 解析失败：{e}")

    requirements = ""
    req_path = os.path.join(in_dir, "requirements.txt")
    if os.path.exists(req_path):
        try:
            with open(req_path, "r", encoding="utf-8") as f:
                requirements = f.read().strip()
        except Exception:
            requirements = ""

    # 解析额外约束中的教师指定（如 T001教C01语文）
    teacher_constraints = _parse_teacher_constraints(requirements, teachers, classes)
    if is_new_format and teacher_constraints:
        courses = _generate_courses_new(classes, teachers, standards, assignments, warnings_list, teacher_constraints)
        warnings_list.append(f"已从额外约束中解析 {len(teacher_constraints)} 条教师指定。")

    p = Problem(classes=classes, teachers=teachers, courses=courses,
                requirements=requirements, config=cfg, warnings=warnings_list)
    _post_process(p)
    return p


def _post_process(p):
    cfg = p.config
    grid = cfg.grid_size()

    if cfg.consecutive.only_core_subjects:
        core = set(cfg.consecutive.subjects)
        for c in p.courses:
            if c.subject not in core and c.block_len > 0:
                p.warnings.append(f"班级 {c.class_id} 的「{c.subject}」连堂已自动取消")
                c.block_len = 0

    alt = [c for c in p.courses if c.week_mode not in ("每周", "", None)]
    if alt:
        p.warnings.append(f"检测到 {len(alt)} 条单双周课程，暂按每周排课")

    if cfg.solver.self_study_fill:
        for ci in p.classes:
            has_ss = any(c.subject == "自习" and c.class_id == ci.id for c in p.courses)
            if has_ss:
                continue
            used = sum(c.weekly for c in p.courses if c.class_id == ci.id)
            rest = grid - used
            if rest > 0:
                p.courses.append(CourseReq(
                    class_id=ci.id, subject="自习", teacher_id=None,
                    weekly=rest, block_len=0, week_mode="每周"))
