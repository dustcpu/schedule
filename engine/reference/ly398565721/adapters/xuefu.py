from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from core.models import LessonRequirement, ScheduleInput


SUBJECT_COLUMNS = {
    5: "语文", 6: "数学", 7: "英语", 8: "物理", 9: "化学", 10: "生物",
    11: "政治", 12: "历史", 13: "地理", 14: "体育", 15: "信息技术",
}
DAYS = ("Monday", "Tuesday", "Wednesday", "Thursday", "Friday")


def load_xuefu_appointment(workbook_path: Path) -> ScheduleInput:
    """Read the established 学府高一聘任表 without altering the source workbook."""
    book = load_workbook(workbook_path, read_only=True, data_only=True)
    if "高一教师聘任表" not in book.sheetnames:
        raise ValueError("未找到工作表：高一教师聘任表")
    sheet = book["高一教师聘任表"]
    rows = tuple(range(4, 23))
    classes = tuple(str(sheet.cell(row, 1).value).strip() for row in rows)
    if any(not value or value == "None" for value in classes):
        raise ValueError("学府聘任表A列4—22行必须均为班级编号")
    lessons: list[LessonRequirement] = []
    for row, class_id in zip(rows, classes, strict=True):
        for column, subject in SUBJECT_COLUMNS.items():
            teacher = sheet.cell(row, column).value
            hours = sheet.cell(3, column).value
            if not teacher or hours is None:
                raise ValueError(f"{class_id}班{subject}缺少教师或周课时")
            weekly_hours = int(hours)
            if subject == "生物":
                weekly_hours = 3 if class_id in {"1", "2"} else 2
            lessons.append(LessonRequirement(class_id, subject, str(teacher).strip(), weekly_hours))
        lessons.append(LessonRequirement(class_id, "艺术（单美术/双音乐）", "艺术轮换", 1, resource="艺术轮换"))
    blocked = {
        class_id: frozenset(
            {("Wednesday", 7), ("Wednesday", 8), ("Friday", 8)}
            | ({("Friday", 7)} if class_id not in {"1", "2"} else set())
        )
        for class_id in classes
    }
    return ScheduleInput(classes, DAYS, tuple(range(1, 9)), tuple(lessons), blocked_slots=blocked)
