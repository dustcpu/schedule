import json
import os
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
NODE = os.environ.get("CODEX_NODE", shutil.which("node") or "node")
pytestmark = pytest.mark.skipif(
    not (ROOT / "node_modules" / "@oai" / "artifact-tool").exists(),
    reason="Excel export test requires Codex artifact-tool runtime; see README setup.",
)


def test_export_contract_has_exactly_four_workbooks_and_name_boundaries(tmp_path: Path):
    schedule_path = tmp_path / "schedule.json"
    duty_path = tmp_path / "duties.json"
    output_dir = tmp_path / "outputs"
    schedule_path.write_text(json.dumps({"status": "FEASIBLE", "lessons": [
        {"class_id": "1", "subject": "语文", "teacher": "Teacher-A", "day": "Monday", "period": 1},
        {"class_id": "1", "subject": "数学", "teacher": "Teacher-B", "day": "Tuesday", "period": 2},
    ]}, ensure_ascii=False), encoding="utf-8")
    duty_path.write_text(json.dumps({"status": "FEASIBLE", "assignments": [
        {"day": "Monday", "class_id": "1", "duty_type": "early", "teacher": "Teacher-A"},
        {"day": "Monday", "class_id": "1", "duty_type": "quiet", "teacher": "Teacher-B"},
        {"day": "Sunday", "class_id": "1", "duty_type": "evening", "teacher": "Teacher-A"},
    ]}, ensure_ascii=False), encoding="utf-8")
    subprocess.run([str(NODE), str(ROOT / "export" / "export_workbooks.mjs"), str(schedule_path), str(duty_path), str(output_dir)], check=True, cwd=ROOT)
    paths = sorted(output_dir.glob("*.xlsx"))
    assert [path.name for path in paths] == ["总表学科版.xlsx", "教师个人课表.xlsx", "早读静校晚自习值班表.xlsx", "班级课表.xlsx"]
    master = load_workbook(output_dir / "总表学科版.xlsx", data_only=True)["总课表"]
    assert master["B4"].value == "语文"
    assert master["B5"].value == "语文"
    duty = load_workbook(output_dir / "早读静校晚自习值班表.xlsx", data_only=True)["值班表"]
    assert [duty.cell(3, column).value for column in range(1, 5)] == ["日期", "班级", "值班类型", "教师姓名"]
    assert duty.max_column == 4
