from core.formal_solver import solve_formal
from core.models import LessonRequirement, ScheduleInput, ScheduledLesson, SolvedSchedule
from core.profile import default_profile
from core.validate_formal import validate_formal
from adapters.xuefu import load_xuefu_appointment
from pathlib import Path
from openpyxl import Workbook


def confirmed_profile():
    values = {
        "F04": {"days": ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"], "periods": 8},
        "F05": [], "F17": {}, "F18": {}, "F19": [],
        "D01": "default", "D13": {}, "D14": False, "D15": "one_to_three_prefer_two", "D16": [],
    }
    profile = default_profile()
    for rule_id, value in values.items():
        profile = profile.with_confirmation(rule_id, value)
    return profile


def two_class_input():
    return ScheduleInput(
        classes=("1", "2"),
        days=("Monday", "Tuesday", "Wednesday", "Thursday", "Friday"),
        periods=tuple(range(1, 9)),
        lessons=(
            LessonRequirement("1", "数学", "T", 3),
            LessonRequirement("2", "数学", "T", 1),
        ),
    )


def test_two_class_teacher_cannot_teach_three_lessons_to_same_class_on_one_day():
    solved = SolvedSchedule(
        status="FEASIBLE",
        lessons=(
            ScheduledLesson("1", "数学", "T", "Monday", 1),
            ScheduledLesson("1", "数学", "T", "Monday", 2),
            ScheduledLesson("1", "数学", "T", "Monday", 3),
            ScheduledLesson("2", "数学", "T", "Tuesday", 1),
        ),
    )
    report = validate_formal(two_class_input(), solved, confirmed_profile())
    assert "TWO_CLASS_SAME_CLASS_THREE" in {failure.code for failure in report.failures}


def test_half_day_infeasibility_returns_minimum_named_exception_candidates():
    schedule_input = ScheduleInput(
        classes=("1",), days=("Monday", "Tuesday", "Wednesday", "Thursday", "Friday"),
        periods=tuple(range(1, 9)), lessons=(LessonRequirement("1", "语文", "T", 18),),
        blocked_slots={"1": frozenset(("Monday", period) for period in range(2, 9) if period != 5)},
    )
    outcome = solve_formal(schedule_input, confirmed_profile(), time_limit_seconds=3)
    assert outcome.status == "INFEASIBLE_PENDING_RELAXATION"
    assert len(outcome.relaxation_candidates) == 1
    assert outcome.relaxation_candidates[0][0] == "T"


def test_solver_enforces_pe_daily_limit_and_teacher_collision():
    schedule_input = ScheduleInput(
        classes=("1", "2"),
        days=("Monday", "Tuesday"), periods=(1, 2, 3, 4),
        lessons=(
            LessonRequirement("1", "体育", "PE", 2),
            LessonRequirement("2", "数学", "M", 2),
        ),
    )
    outcome = solve_formal(schedule_input, confirmed_profile(), time_limit_seconds=3)
    assert outcome.status in {"OPTIMAL", "FEASIBLE"}
    assert validate_formal(schedule_input, outcome, confirmed_profile()).ok


def test_xuefu_adapter_reads_a_fictional_appointment_table(tmp_path: Path):
    source = tmp_path / "fictional_xuefu_appointment.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "高一教师聘任表"
    weekly_hours = {5: 5, 6: 5, 7: 5, 8: 4, 9: 4, 10: 2, 11: 2, 12: 2, 13: 2, 14: 2, 15: 1}
    for column, hours in weekly_hours.items():
        sheet.cell(3, column).value = hours
    for row in range(4, 23):
        sheet.cell(row, 1).value = row - 3
        for column in weekly_hours:
            sheet.cell(row, column).value = f"Teacher-{column}"
    workbook.save(source)
    result = load_xuefu_appointment(source)
    assert len(result.classes) == 19
    assert len(result.lessons) == 19 * 12
    assert result.weekly_hours[("1", "生物")] == 3
    assert result.weekly_hours[("3", "生物")] == 2
