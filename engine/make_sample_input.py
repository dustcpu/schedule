# -*- coding: utf-8 -*-
"""生成最小可用测试 input.xlsx（协议 §10 要求算法团队交付一份供外壳联调）。

用法: python make_sample_input.py <输出路径.xlsx>

说明：数据全部为占位符（团队约定不写真实教师/学生姓名），教师用 T01/T02… 编号。
"""
import sys

from openpyxl import Workbook


def build(path: str, n_classes: int = 4) -> None:
    wb = Workbook()

    # ---- 班级 ----
    ws = wb.active
    ws.title = "班级"
    ws.append(["班级ID", "班级名称", "年级", "科类"])
    for i in range(1, n_classes + 1):
        ws.append([f"C{i:02d}", f"高二({i})班", "高二", "理科"])

    # ---- 教师 ----
    ws = wb.create_sheet("教师")
    ws.append(["教师ID", "教师姓名", "任教学科"])
    subjects = ["语文", "数学", "英语", "物理", "化学", "生物", "体育"]
    tid = 1
    for s in subjects:
        # 每个学科两位教师（可带多个班，便于验证教师冲突约束）
        for _ in range(2):
            ws.append([f"T{tid:02d}", f"教师{tid:02d}", s])
            tid += 1

    # ---- 课程 ----
    # 说明：连堂节数 2 表示连堂；自习留空由内核自动补齐
    ws = wb.create_sheet("课程")
    ws.append(["班级ID", "学科", "教师ID", "周课时", "连堂节数", "单双周"])

    # 主科课时较高，副科较低；总和需 <= 40（8节×5天）
    plan = [
        ("语文", 6, 2), ("数学", 6, 2), ("英语", 6, 2),
        ("物理", 5, 0), ("化学", 4, 0), ("生物", 3, 0), ("体育", 2, 0),
    ]
    # 教师分配：每班按序取该学科的第 1 位教师，相邻班错开用第 2 位
    teacher_index = {}
    for si, (s, _, _) in enumerate(plan):
        teacher_index[s] = [si * 2 + 1, si * 2 + 2]

    # 固定课学科：对应外壳默认固定课（周一第1节班会/周一第8节研究性学习/周四第8节校本课）
    # 教师ID 留空 = 无固定教师，不参与教师冲突（如班会由各班主任在本班上）
    fixed_plan = [("班会", 1), ("研究性学习", 1), ("校本课", 1)]

    for ci in range(1, n_classes + 1):
        cid = f"C{ci:02d}"
        for s, weekly, block in plan:
            tno = teacher_index[s][(ci - 1) % 2]
            ws.append([cid, s, f"T{tno:02d}", weekly, block, "每周"])
        for s, weekly in fixed_plan:
            ws.append([cid, s, "", weekly, 0, "每周"])

    wb.save(path)
    print(f"已生成 {path}")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "sample_input.xlsx"
    n = int(sys.argv[2]) if len(sys.argv) > 2 else 4
    build(out, n)
