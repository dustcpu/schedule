# -*- coding: utf-8 -*-
"""输入文件校验：支持新格式（教师/班级/课时标准/任课安排）和旧格式。"""
import json
import sys
import os


def validate(path):
    errors = []
    warnings = []
    stats = {}

    if not os.path.exists(path):
        return {"valid": False, "errors": ["文件不存在"], "warnings": [], "stats": {}}
    if not path.lower().endswith((".xlsx", ".xlsm")):
        errors.append("文件格式不是 .xlsx")

    try:
        from openpyxl import load_workbook
    except ImportError:
        return {"valid": False, "errors": ["缺少 openpyxl 库"], "warnings": [], "stats": {}}

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:
        return {"valid": False, "errors": [f"无法打开文件: {e}"], "warnings": [], "stats": {}}

    sheet_names = wb.sheetnames

    def find_sheet(names):
        for n in names:
            for sn in sheet_names:
                if sn.strip() == n:
                    return sn
        return None

    def read_rows(ws):
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], []
        headers = [str(h).strip() if h else "" for h in rows[0]]
        data = []
        for r in rows[1:]:
            if r is None or all(c is None or str(c).strip() == "" for c in r):
                continue
            d = {}
            for i, h in enumerate(headers):
                if h:
                    d[h] = r[i] if i < len(r) else None
            data.append(d)
        return headers, data

    period_sheet = find_sheet(["课时标准", "课时"])
    is_new_format = period_sheet is not None
    stats["format"] = "新格式（课时标准）" if is_new_format else "旧格式（课程表）"

    teacher_sheet = find_sheet(["教师", "教师表"])
    teachers = {}
    if teacher_sheet is None:
        errors.append("缺少「教师」工作表")
    else:
        headers, rows = read_rows(wb[teacher_sheet])
        for col in ["教师ID", "任教学科", "教师姓名"]:
            if col not in headers:
                errors.append(f"「教师」表缺少列：{col}")
        if not errors:
            for r in rows:
                tid = str(r.get("教师ID") or "").strip()
                if not tid:
                    continue
                teachers[tid] = {
                    "name": str(r.get("教师姓名") or "").strip() or tid,
                    "subject": str(r.get("任教学科") or "").strip(),
                }
            stats["teachers"] = len(teachers)
            if len(teachers) == 0:
                errors.append("「教师」表没有有效数据")

    class_sheet = find_sheet(["班级", "班级表"])
    classes = {}
    if class_sheet is None:
        errors.append("缺少「班级」工作表")
    else:
        headers, rows = read_rows(wb[class_sheet])
        for col in ["班级ID", "班级名称", "年级"]:
            if col not in headers:
                errors.append(f"「班级」表缺少列：{col}")
        if not errors:
            for r in rows:
                cid = str(r.get("班级ID") or "").strip()
                if not cid:
                    continue
                classes[cid] = {
                    "name": str(r.get("班级名称") or "").strip() or cid,
                    "elective": str(r.get("选科") or "").strip(),
                }
            stats["classes"] = len(classes)
            if len(classes) == 0:
                errors.append("「班级」表没有有效数据")

    if errors:
        wb.close()
        return {"valid": False, "errors": errors, "warnings": warnings, "stats": stats}

    if is_new_format:
        headers, rows = read_rows(wb[period_sheet])
        for col in ["学科", "选考周课时", "非选考周课时"]:
            if col not in headers:
                errors.append(f"「课时标准」表缺少列：{col}")
        if not errors:
            standards = {}
            for r in rows:
                subj = str(r.get("学科") or "").strip()
                if not subj:
                    continue
                try:
                    elec = int(float(str(r.get("选考周课时") or 0)))
                    non_elec = int(float(str(r.get("非选考周课时") or 0)))
                except (TypeError, ValueError):
                    errors.append(f"「课时标准」表中「{subj}」的课时格式错误")
                    continue
                standards[subj] = {"elective": elec, "non_elective": non_elec}
            stats["subjects"] = len(standards)
            if len(standards) == 0:
                errors.append("「课时标准」表没有有效数据")

            subj_map = {"物": "物理", "化": "化学", "生": "生物", "政": "政治", "史": "历史", "地": "地理"}
            for cid, ci in classes.items():
                elective_set = set()
                if ci["elective"]:
                    for ch in ci["elective"]:
                        if ch in subj_map:
                            elective_set.add(subj_map[ch])
                total = 0
                for subj, s in standards.items():
                    weekly = s["elective"] if subj in elective_set else s["non_elective"]
                    total += weekly
                if total > 40:
                    errors.append(f"班级 {ci['name']} 周课时合计 {total} 超过可用格数 40")
                elif total == 40:
                    warnings.append(f"班级 {ci['name']} 周课时刚好 40，没有自习时间")

        assign_sheet = find_sheet(["任课安排", "任课"])
        if assign_sheet:
            headers, rows = read_rows(wb[assign_sheet])
            unknown_t = set()
            unknown_c = set()
            for r in rows:
                tid = str(r.get("教师ID") or "").strip()
                cid = str(r.get("班级ID") or "").strip()
                if tid and tid not in teachers:
                    unknown_t.add(tid)
                if cid and cid not in classes:
                    unknown_c.add(cid)
            if unknown_t:
                warnings.append(f"任课安排中有 {len(unknown_t)} 个教师ID在教师表中不存在")
            if unknown_c:
                warnings.append(f"任课安排中有 {len(unknown_c)} 个班级ID在班级表中不存在")
            stats["assignments"] = len(rows)
        else:
            warnings.append("没有「任课安排」表，将按学科自动分配教师")
    else:
        course_sheet = find_sheet(["课程", "课程表", "课时"])
        if course_sheet is None:
            errors.append("缺少「课程」工作表")
        else:
            headers, rows = read_rows(wb[course_sheet])
            for col in ["班级ID", "学科", "周课时"]:
                if col not in headers:
                    errors.append(f"「课程」表缺少列：{col}")
            if not errors:
                unknown_cids = set()
                for r in rows:
                    cid = str(r.get("班级ID") or "").strip()
                    subj = str(r.get("学科") or "").strip()
                    if cid and cid not in classes:
                        unknown_cids.add(cid)
                    try:
                        weekly = int(float(str(r.get("周课时") or 0)))
                        if weekly <= 0:
                            errors.append(f"班级 {cid} 的「{subj}」周课时应为正整数")
                    except (TypeError, ValueError):
                        errors.append(f"班级 {cid} 的「{subj}」周课时格式错误")
                if unknown_cids:
                    errors.append(f"课程表中有 {len(unknown_cids)} 个班级ID不存在")
                stats["courses"] = len(rows)

    wb.close()
    return {"valid": len(errors) == 0, "errors": errors, "warnings": warnings, "stats": stats}


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"valid": False, "errors": ["缺少文件路径"], "warnings": [], "stats": {}}))
        sys.exit(1)
    result = validate(sys.argv[1])
    print(json.dumps(result, ensure_ascii=False))
