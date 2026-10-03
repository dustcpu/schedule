# -*- coding: utf-8 -*-
"""nl2req —— 外置「自然语言 → 结构化约束」转换器（命令行，离线）。

用法：
  uv run --with openpyxl python nl2req.py --text 语料.txt --xlsx input.xlsx \
      [--hard-limits hard_limits.json] [--out structured_requirements.json]
  python nl2req.py --line "张老师不想周二上连堂" --xlsx input.xlsx
  python nl2req.py --validate ai返回.json --xlsx input.xlsx [--out 清洗后.json]

产出：
  1. structured_requirements.json —— 放进引擎输入目录即可生效（可选通道）；
  2. 一段 ```json 围栏块 —— 整段粘贴进软件「特殊要求」框即可生效（兼容通道）；
  3. 认不出的行 → 生成可外发免费 AI 的提示词（姓名已替换为编号，脱敏），
     AI 返回的 JSON 用 --validate 校验后再用。

定位（docs/额外约束处理方案评估.md 方案3 + 2026-10-03 定案）：
  本工具不进软件发布物、不进引擎进程、全程不联网；词典命中优先，
  模型只作为人工外发的兜底通道，且输出必须经本工具校验。
"""
import argparse
import json
import os
import re
import sys

_HERE = os.path.dirname(os.path.abspath(__file__))
_REPO_ROOT = os.path.dirname(_HERE)
_ENGINE_DIR = os.path.join(_REPO_ROOT, "engine")
if os.path.isdir(_ENGINE_DIR):
    sys.path.insert(0, _ENGINE_DIR)

from 词典 import (  # noqa: E402
    SCHEMA_VERSION, CONSTRAINT_TYPES, ParseContext, clean_line,
    parse_line, describe_constraint, build_ai_prompt,
)

try:
    from scheduler.data.load import (  # noqa: E402  复用引擎的加载器，保证列名/别名口径一致
        _load_classes, _load_teachers, _resolve_id, normalize_subject,
    )
    from scheduler.config import load_from_hard_limits  # noqa: E402
    _ENGINE_OK = True
except Exception:  # 引擎不在旁时仍可独立跑（学科归一化降级为原样返回）
    _ENGINE_OK = False

    def normalize_subject(name):
        return name

    def _resolve_id(token, ids):
        return token if token in ids else None


if _ENGINE_OK:
    from openpyxl import load_workbook
else:
    try:
        from openpyxl import load_workbook
    except ImportError:
        load_workbook = None


# ---------------------------------------------------------------- 输入上下文

class XlsxContext(ParseContext):
    """用 input.xlsx 的教师/班级表做实体解析（全部本地，不上传）。"""

    def __init__(self, classes, teachers, cfg):
        self.classes = {c.id: c for c in classes}
        self.teachers = teachers
        self.morning = cfg.schedule.morning_periods
        self.per_day = cfg.schedule.periods_per_day
        consec = cfg.consecutive
        self._day_of = {s: consec.day_of(s) for s in consec.subjects}
        self._subjects = sorted({t.subject for t in teachers.values() if t.subject},
                                key=len, reverse=True)
        self._name_to_tid = {}
        for tid, t in teachers.items():
            if t.name:
                self._name_to_tid.setdefault(t.name, tid)

    # ---- 教师 ----
    def resolve_teacher(self, token):
        token = (token or "").strip()
        if not token:
            return None, None
        if token in self.teachers:
            return token, None
        hit = _resolve_id(token, list(self.teachers)) if _ENGINE_OK else None
        if hit:
            return hit, None
        # 姓名：精确 → 去「老师」后缀 → 姓氏匹配
        name = token[:-2] if token.endswith("老师") else token
        if name in self._name_to_tid:
            return self._name_to_tid[name], None
        same = [tid for tid, t in self.teachers.items()
                if t.name and (t.name == name or t.name.startswith(name))]
        if len(same) == 1:
            return same[0], None
        if len(same) > 1:
            show = "、".join(self.teacher_display(t) for t in same[:5])
            return None, f"「{token}」对应多位教师（{show}），请写全名或教师ID"
        near = f"，表里有「{next(iter(self.teachers))}」等" if self.teachers else ""
        return None, f"「{token}」不在「教师」表里{near}，请照抄表里的姓名或教师ID"

    # ---- 班级 ----
    def resolve_class(self, token):
        token = (token or "").strip()
        if not token:
            return None, None
        names = {c.id: (c.name or c.id) for c in self.classes.values()}
        if token in self.classes:
            return token, None
        hit = _resolve_id(token, list(self.classes)) if _ENGINE_OK else None
        if hit:
            return hit, None
        exact = [cid for cid, nm in names.items() if nm == token]
        if len(exact) == 1:
            return exact[0], None
        part = [cid for cid, nm in names.items() if token in nm]
        if len(part) == 1:
            return part[0], None
        if len(part) > 1:
            show = "、".join(sorted(names[c] for c in part)[:5])
            return None, f"「{token}」对应多个班级（{show}），请写完整班级名或班级ID"
        return None, f"「{token}」不在「班级」表里，请照抄表里的班级名称或班级ID"

    # ---- 学科 ----
    def match_subject(self, subj_raw, tid):
        subj_raw = (subj_raw or "").strip()
        if not subj_raw:
            return None
        subj_raw = normalize_subject(subj_raw)
        t = self.teachers.get(tid)
        if t and t.subject and (t.subject in subj_raw or subj_raw == t.subject):
            return t.subject
        for s in self._subjects:
            if s and s in subj_raw:
                return s
        return None

    # ---- 展示与口径 ----
    def teacher_display(self, tid):
        t = self.teachers.get(tid)
        if t and t.name and t.name != tid:
            return f"{t.name}（{tid}）"
        return tid

    def teacher_info(self, tid):
        t = self.teachers.get(tid)
        return {"name": t.name, "subject": t.subject} if t else None

    def double_day_of(self, subject):
        return self._day_of.get(subject)


def _load_ctx(xlsx_path, hard_limits_path):
    if load_workbook is None:
        raise SystemExit("缺少 openpyxl：请用 uv run --with openpyxl python nl2req.py …")
    if not os.path.exists(xlsx_path):
        raise SystemExit(f"找不到 {xlsx_path}")
    try:
        wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    except Exception as e:
        raise SystemExit(f"打不开「{xlsx_path}」：{e}。请先关闭正在使用它的 Excel。")

    def _find(names):
        norm = {str(s).strip().replace("　", ""): s for s in wb.sheetnames}
        for n in names:
            if n in norm:
                return wb[norm[n]]
        return None

    classes = _load_classes(_find(["班级", "班级表", "classes"]))
    teachers = _load_teachers(_find(["教师", "教师表", "teachers"]))
    wb.close()
    if not teachers:
        raise SystemExit("「教师」sheet 为空，无法解析教师姓名")

    cfg = None
    if _ENGINE_OK:
        from scheduler.config import Config
        cfg = Config()
        if hard_limits_path and os.path.exists(hard_limits_path):
            try:
                with open(hard_limits_path, "r", encoding="utf-8-sig") as f:
                    cfg = load_from_hard_limits(json.load(f))
            except Exception as e:
                print(f"⚠ hard_limits.json 解析失败（{e}），改用引擎默认口径", file=sys.stderr)
    if cfg is None:  # 引擎不可用时的兜底口径（与引擎默认一致）
        from types import SimpleNamespace
        cfg = SimpleNamespace(
            schedule=SimpleNamespace(morning_periods=5, periods_per_day=8),
            consecutive=SimpleNamespace(
                subjects=["语文", "数学", "英语"],
                day_of={"语文": 3, "数学": 2, "英语": 4}.get),
        )
        cfg.consecutive.day_of = (lambda s: {"语文": 3, "数学": 2, "英语": 4}.get(s))
    return XlsxContext(classes, teachers, cfg)


# ---------------------------------------------------------------- 主流程

def _dedup(constraints):
    out, seen = [], set()
    for c in constraints:
        key = json.dumps(c, sort_keys=True, ensure_ascii=False)
        if key not in seen:
            seen.add(key)
            out.append(c)
    return out


def _desensitize(line, ctx):
    """外发前把行内真实姓名/班级名替换成编号（隐私红线：姓名不出本机）。"""
    out = line
    for tid, t in sorted(ctx.teachers.items(),
                         key=lambda kv: -len(kv[1].name or "")):
        if t.name and len(t.name) >= 2 and t.name in out:
            out = out.replace(t.name, tid)
    for cid, c in ctx.classes.items():
        if c.name and len(c.name) >= 2 and c.name in out:
            out = out.replace(c.name, cid)
    return out


def run_convert(args):
    ctx = _load_ctx(args.xlsx, args.hard_limits)

    if args.line:
        lines = [args.line]
    else:
        if not args.text or not os.path.exists(args.text):
            raise SystemExit(f"找不到 {args.text}（或改用 --line \"一句话\"）")
        with open(args.text, "r", encoding="utf-8-sig") as f:
            lines = f.read().splitlines()

    applied, unknown, empties = [], [], 0
    print(f"共 {len([l for l in lines if l.strip()])} 行，逐行解析：\n")
    for i, raw in enumerate(lines, 1):
        if not raw.strip():
            continue
        r = parse_line(raw, ctx)
        if r["status"] == "empty":
            empties += 1
            continue
        no = f"{i:>3}"
        if r["status"] == "ok":
            applied.extend(r["constraints"])
            print(f"✓ {no} 「{r['raw']}」")
            print(f"      我理解成：{r['echo']}")
            for note in r["notes"]:
                print(f"      ⚠ {note}")
        else:
            unknown.append((i, r["raw"], r["reason"]))
            print(f"× {no} 「{r['raw']}」")
            print(f"      认不出：{r['reason']}")

    applied = _dedup(applied)
    print(f"\n小结：命中 {len(applied)} 条约束（已去重），认不出 {len(unknown)} 行。")

    # ---- 产出 1：结构化 JSON ----
    payload = {"version": SCHEMA_VERSION, "constraints": applied}
    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
            f.write("\n")
        print(f"已写出 {args.out}（放进引擎输入目录即可生效）")

    # ---- 产出 2：可粘贴的围栏块 ----
    if applied:
        block = json.dumps(payload, ensure_ascii=False, indent=2)
        print("\n把下面整段（含 ``` 行）粘贴进软件「特殊要求」框即可生效：\n")
        print("```json")
        print(block)
        print("```")

    # ---- 产出 3：认不出的行 → AI 提示词（姓名已脱敏） ----
    if unknown and not args.no_ai_hint:
        sanitized = [f"{_desensitize(ln, ctx)}" for _i, ln, _r in unknown]
        print("\n" + "=" * 60)
        print("以下行词典认不出。可把【提示词】整段发给任意免费 AI，"
              "把返回的 JSON 存成文件后用 --validate 校验：\n")
        print(build_ai_prompt(sanitized))
    elif unknown:
        print("\n有认不出的行（已按 --no-ai-hint 跳过 AI 提示词生成）。")
    return 0 if not unknown else 1


VALID_FIELD_CHECKS = {
    "assign_teacher": ("teacher", "class", "subject"),
    "teacher_unavailable": ("teacher", "days"),
    "teacher_no_double_day": ("teacher", "days"),
    "class_unavailable": ("class", "day"),
    "teacher_max_classes": ("teacher", "max"),
}


def run_validate(args):
    """校验外部 AI 返回的 JSON：schema 结构 + 实体存在性。逐条独立，坏一条不牵连。"""
    ctx = _load_ctx(args.xlsx, args.hard_limits)
    try:
        with open(args.validate, "r", encoding="utf-8-sig") as f:
            data = json.load(f)
    except Exception as e:
        raise SystemExit(f"读不上 JSON：{e}（AI 的输出要存成纯 JSON 文件，"
                         f"别把解释文字一起贴进来）")
    if not isinstance(data, dict) or "constraints" not in data:
        raise SystemExit("JSON 里没有 constraints 字段，不是本工具约定的格式")
    version = data.get("version", SCHEMA_VERSION)
    if version != SCHEMA_VERSION:
        print(f"⚠ version={version}，本工具按 v{SCHEMA_VERSION} 校验")

    ok_list, bad = [], []
    for i, c in enumerate(data.get("constraints", []), 1):
        err = _validate_one(c, ctx)
        if err:
            bad.append((i, c, err))
        else:
            ok_list.append(c)

    for c in ok_list:
        print(f"✓ {describe_constraint(c, ctx)}")
    for i, c, err in bad:
        print(f"× 第{i}条 {json.dumps(c, ensure_ascii=False)}")
        print(f"      不合法：{err}")
    for ln in (data.get("unrecognized") or []):
        print(f"· AI 也认不出：{ln}")

    print(f"\n校验结果：{len(ok_list)} 条可用，{len(bad)} 条被拒。")
    if args.out:
        payload = {"version": SCHEMA_VERSION, "constraints": ok_list}
        with open(args.out, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
            f.write("\n")
        print(f"已把可用约束写出 {args.out}")
    return 0 if not bad else 2


def _validate_one(c, ctx):
    if not isinstance(c, dict):
        return "不是对象"
    t = c.get("type")
    if t not in CONSTRAINT_TYPES:
        return (f"未知类型「{t}」。只支持：{'、'.join(CONSTRAINT_TYPES)}"
                f"（宁可弃权，不要猜）")
    for field in VALID_FIELD_CHECKS[t]:
        if c.get(field) in (None, "", []):
            return f"缺少必填字段「{field}」"
    days = c.get("days", [])
    if isinstance(days, int):
        days = [days]
    for d in (days if isinstance(days, list) else []):
        if not isinstance(d, int) or not (1 <= d <= 5):
            return f"days 里的「{d}」不是 1-5（周一到周五）"
    day = c.get("day")
    if day is not None and (not isinstance(day, int) or not (1 <= day <= 5)):
        return f"day「{day}」不是 1-5（周一到周五）"
    periods = c.get("periods", [])
    if isinstance(periods, int):
        periods = [periods]
    for p in (periods if isinstance(periods, list) else []):
        if not isinstance(p, int) or not (1 <= p <= ctx.per_day):
            return f"periods 里的「{p}」超出 1-{ctx.per_day} 节"
    if t in ("assign_teacher", "teacher_unavailable",
             "teacher_no_double_day", "teacher_max_classes"):
        tid, err = ctx.resolve_teacher(str(c.get("teacher", "")))
        if not tid:
            return err or f"教师「{c.get('teacher')}」不在「教师」表里"
    if t == "assign_teacher":
        if not ctx.match_subject(str(c.get("subject", "")), tid):
            return f"学科「{c.get('subject')}」认不出（要照「教师」sheet 的任教学科写）"
    if t == "teacher_max_classes":
        try:
            if int(c.get("max")) < 1:
                return "max 必须 ≥ 1"
        except (TypeError, ValueError):
            return "max 必须是数字"
    return None


def main():
    ap = argparse.ArgumentParser(
        description="自然语言 → 结构化排课约束（离线转换器）")
    ap.add_argument("--text", help="待解析的文本文件（一行一条要求）")
    ap.add_argument("--line", help="直接解析一句话（调试用）")
    ap.add_argument("--xlsx", help="input.xlsx 路径（提供教师/班级姓名对照）")
    ap.add_argument("--hard-limits", help="可选：hard_limits.json（读取实际连堂日/作息口径）")
    ap.add_argument("--out", help="输出 JSON 文件路径（默认只打印）")
    ap.add_argument("--validate", help="校验模式：外部 AI 返回的 JSON 文件")
    ap.add_argument("--no-ai-hint", action="store_true",
                    help="不为认不出的行生成 AI 提示词")
    args = ap.parse_args()

    if not args.xlsx:
        ap.error("必须给 --xlsx（教师/班级姓名对照全靠它，且只在本地读取）")
    if args.validate:
        sys.exit(run_validate(args))
    if args.text or args.line:
        sys.exit(run_convert(args))
    ap.error("给 --text <文件> 或 --line \"一句话\"；校验外部 AI 返回用 --validate")


if __name__ == "__main__":
    main()
