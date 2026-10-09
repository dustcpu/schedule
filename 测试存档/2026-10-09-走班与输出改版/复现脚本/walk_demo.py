# -*- coding: utf-8 -*-
"""走班：端到端复现脚本（2026-10-09）。

做三件事：
  1. 用仓库自带的合成数据（test_input_25_new，占位命名）造一份"含走班"的输入；
  2. 跑引擎（源码版）排课；
  3. 用 verify_hard 复核交付课表，并打印走班格的实际内容。

全程只用合成数据，不涉及任何真实师生信息。
用法：D:\\python\\python.exe walk_demo.py
"""
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
ENGINE = os.path.join(REPO, "engine")
SRC = os.path.join(ENGINE, "test_input_25_new", "input.xlsx")

# 参与走班的班 + 走班科目，按合成数据的选科来配：
#   C01/C04 = 政史地（政治、地理都有）
#   C23     = 物生政（只有政治）
PARTICIPANTS = ["C01", "C04", "C23"]


def build_sample(dst_in):
    """在 dst_in 里生成 input.xlsx（合成数据 + 两张走班 sheet）。"""
    from openpyxl import load_workbook

    # 找政治/地理老师（占位命名）
    wb0 = load_workbook(SRC, data_only=True, read_only=True)
    rows = list(wb0["教师"].iter_rows(values_only=True))
    hdr = [str(x).strip() if x else "" for x in rows[0]]
    i_id, i_sub = hdr.index("教师ID"), hdr.index("任教学科")
    pol, geo = [], []
    for r in rows[1:]:
        if not r or not r[i_id]:
            continue
        s = str(r[i_sub] or "").strip()
        if s == "政治":
            pol.append(r[i_id])
        elif s == "地理":
            geo.append(r[i_id])
    wb0.close()
    if len(pol) < 2 or len(geo) < 2:
        raise SystemExit("合成数据里的政治/地理老师不够，无法造样本")

    wb = load_workbook(SRC)
    ws = wb.create_sheet("走班时间")
    ws.append(["独立时间ID", "周次数", "连排节数", "参与班级"])
    ws.append(["T1", 2, 2, ",".join(PARTICIPANTS)])

    ws2 = wb.create_sheet("走班教学班")
    ws2.append(["教学班ID", "独立时间", "学科", "来源班级", "人数", "教室", "任课教师"])
    for row in [
        ["政治走班1", "T1", "政治", "C01", 20, "C01教室", pol[0]],
        ["政治走班1", "T1", "", "C23", 22, "", ""],          # 同一教学班的第二个来源
        ["政治走班2", "T1", "政治", "C04", 24, "C04教室", pol[1]],
        ["地理走班1", "T1", "地理", "C01", 25, "走班教室1", geo[0]],
        ["地理走班1", "T1", "", "C04", 23, "", ""],
    ]:
        ws2.append(row)
    wb.save(os.path.join(dst_in, "input.xlsx"))

    with io.open(os.path.join(dst_in, "hard_limits.json"), "w", encoding="utf-8") as f:
        json.dump({"solver": {"max_time_seconds": 30, "num_plans": 2, "min_plans": 2}}, f)


def main():
    tmp = tempfile.mkdtemp(prefix="paike_walk_demo_")
    in_dir, out_dir = os.path.join(tmp, "in"), os.path.join(tmp, "out")
    os.makedirs(in_dir)
    os.makedirs(out_dir)
    print("工作目录:", tmp)

    build_sample(in_dir)
    print("样本已生成（含 走班时间 / 走班教学班 两张 sheet）\n")

    print("=== 跑引擎 ===")
    rc = subprocess.run([sys.executable, os.path.join(ENGINE, "engine.py"),
                         in_dir, out_dir, "walk_demo"], cwd=ENGINE).returncode
    print("引擎退出码:", rc)

    st_path = os.path.join(out_dir, "status.json")
    if os.path.exists(st_path):
        st = json.load(io.open(st_path, encoding="utf-8"))
        print("status: code=%s, 方案数=%s" % (st.get("code"), len(st.get("plans", []))))
        for w in st.get("warnings", []):
            if "走班" in w:
                print("  告警:", w)

    print("\n=== 复核交付课表（verify_hard）===")
    subprocess.run([sys.executable, os.path.join(ENGINE, "verify_hard.py"),
                    in_dir, "--xlsx", os.path.join(out_dir, "result.xlsx")], cwd=ENGINE)

    # 打印走班格内容，便于肉眼确认
    xlsx = os.path.join(out_dir, "result.xlsx")
    if os.path.exists(xlsx):
        from openpyxl import load_workbook
        wb = load_workbook(xlsx, data_only=True)
        print("\n=== 走班格（方案1）===")
        ws = wb["方案1"]
        rows = list(ws.iter_rows(values_only=True))
        head = [("" if c is None else str(c)).strip() for c in rows[0]]
        cur_day = ""
        for r in rows[1:]:
            vals = [("" if c is None else str(c)) for c in r]
            if len(vals) < 3:
                continue
            if vals[0].strip():
                cur_day = vals[0].strip()
            for i, v in enumerate(vals):
                if "→" in str(v):
                    print("  %s %s  %s: %s" % (cur_day, vals[1], head[i],
                                               str(v).replace("\n", " ／ ")))
        wb.close()
    print("\n输出目录:", out_dir)


if __name__ == "__main__":
    main()
