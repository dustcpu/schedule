# -*- coding: utf-8 -*-
"""验证用户报的两个问题（只读：不修改仓库任何文件，全部在临时目录里做）。

问题1：课时标准里加一门「信息课」，引擎会不会判定"没有对应教师"？
问题2：特殊要求里那两行，解析器分别怎么处理？
"""
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))   # 仓库根
ENGINE = os.path.join(REPO, "engine")
SRC_XLSX = os.path.join(ENGINE, "test_input_25_new", "input.xlsx")
sys.path.insert(0, ENGINE)

# ============================================================ 问题2：解析器实测
print("=" * 74)
print("【问题2】特殊要求解析器实测")
from scheduler.data.load import (_parse_teacher_constraints, normalize_subject,
                                 TeacherInfo, ClassInfo, NO_TEACHER_SUBJECTS)

TEACHERS = {
    "T30": TeacherInfo(id="T30", name="张三", subject="数学"),
    "T01": TeacherInfo(id="T01", name="李四", subject="语文"),
}
CLASSES = [ClassInfo(id="C01", name="一班"), ClassInfo(id="C05", name="五班")]

LINES = [
    "1.非物理选科班每周安排一节物理，非政治选科班每周安排一节政治",
    "2.教师T30所教授的班级其中一个必须是C05",
    "3.T30教C05数学",                      # 受支持的写法 + 行首编号，作为对照
]
for ln in LINES:
    # ⚠️ 返回值自 2026-10-01（5b03d06）起是 3 元组：(result, unparsed, unsupported)
    res, unparsed, unsupported = _parse_teacher_constraints(ln, TEACHERS, CLASSES)
    if res:
        tag = "✔ 解析成功"
    elif unparsed:
        tag = "△ 语法告警（已提示用户）"
    else:
        tag = "△ 不支持告警（已提示用户）"
    print(f"  「{ln}」")
    print(f"      → {tag}  解析结果={res}")
    print(f"         未解析={unparsed if unparsed else '空'}"
          f"  不支持={unsupported if unsupported else '空'}")

print()
print("  NO_TEACHER_SUBJECTS（不需要教师的学科白名单，2026-10-01 已补 书法/劳动/生涯规划）:")
print("    " + "、".join(sorted(NO_TEACHER_SUBJECTS)))
print("  学科名归一化实测（normalize_subject）：")
for raw in ["信息课", "信息", "计算机", "信息技术", "心理健康", "体育与健康", "语文"]:
    print(f"    {raw:6s} → {normalize_subject(raw)}")

# ============================================================ 问题1：加一门「信息课」
print()
print("=" * 74)
print("【问题1】课时标准里加一门「信息课」（周课时 1），跑一遍引擎")
BASE = os.path.join(tempfile.gettempdir(), "paike_issue_probe")
if os.path.exists(BASE):
    shutil.rmtree(BASE)
inp, out = os.path.join(BASE, "input"), os.path.join(BASE, "output")
os.makedirs(inp)
os.makedirs(out)
shutil.copy(SRC_XLSX, os.path.join(inp, "input.xlsx"))

from openpyxl import load_workbook
wb = load_workbook(os.path.join(inp, "input.xlsx"))
ws = wb["课时标准"]
heads = [str(c.value).strip() if c.value else "" for c in ws[1]]
print("  课时标准表头:", heads)
row = ws.max_row + 1
values = {"学科": "信息课", "选考周课时": 0, "非选考周课时": 1, "连堂节数": 0, "连堂日": 0}
for j, h in enumerate(heads, start=1):
    ws.cell(row=row, column=j, value=values.get(h, ""))
wb.save(os.path.join(inp, "input.xlsx"))
wb.close()

# 教师表里有没有「信息课」这个学科
wb = load_workbook(os.path.join(inp, "input.xlsx"))
ws = wb["教师"]
heads = [str(c.value).strip() if c.value else "" for c in ws[1]]
col = heads.index("任教学科") + 1
subs = set()
for r in range(2, ws.max_row + 1):
    v = str(ws.cell(row=r, column=col).value or "").strip()
    if v:
        subs.add(v)
wb.close()
print("  教师表里出现过的任教学科:", "、".join(sorted(subs)))
print("  有「信息课」教师吗：", "信息课" in subs)

HL = {
    "schedule": {"periods_per_day": 8, "period_minutes": 40, "morning_start": "08:00",
                 "morning_end": "12:20", "afternoon_start": "14:30", "afternoon_end": "16:55",
                 "long_break_after_period": 3, "long_break_minutes": 30,
                 "eye_break_after_period": 6, "eye_break_minutes": 15,
                 "default_break_minutes": 10},
    "fixed_classes": [{"day": 1, "period": 1, "subject": "班会", "note": ""}],
    "classes": {"total_classes": 25, "arts_start": 1, "arts_end": 4, "science_start": 5,
                "science_end": 25, "gaokao_policy": "新高考3+1+2"},
    "consecutive": {"subjects": ["语文", "数学", "英语"], "math_day": 2, "chinese_day": 3,
                    "english_day": 4, "rule": "", "only_core_subjects": True},
    "solver": {"max_time_seconds": 20, "teacher_candidate_k": 1, "num_plans": 1,
               "min_plans": 1, "teacher_decision": False},
    "extra_constraints": "",
}
with open(os.path.join(inp, "requirements.txt"), "w", encoding="utf-8") as f:
    f.write("")
with open(os.path.join(inp, "hard_limits.json"), "w", encoding="utf-8") as f:
    json.dump(HL, f, ensure_ascii=False, indent=2)

# 引擎用当前解释器跑（uv run 或装好 ortools 的 python 均可；不再硬编码机器路径）
p = subprocess.run([sys.executable, os.path.join(ENGINE, "engine.py"), inp, out, "issue1"],
                   cwd=ENGINE, capture_output=True, text=True, encoding="utf-8",
                   errors="replace", env=dict(os.environ, PYTHONIOENCODING="utf-8"))
st = os.path.join(out, "status.json")
if os.path.exists(st):
    d = json.load(open(st, encoding="utf-8"))
    print(f"  >>> code={d['code']}  方案数={len(d['plans'])}")
    print(f"  >>> 提示语: {d['message']}")
else:
    print("  >>> 没有 status.json")
print()
print("  engine.log 里的相关行:")
lg = os.path.join(out, "engine.log")
if os.path.exists(lg):
    for line in open(lg, encoding="utf-8").read().splitlines():
        if "信息" in line or "校验失败" in line or "阶段" in line:
            print("    " + line)
print()
print("完成（临时目录: %s）" % BASE)
