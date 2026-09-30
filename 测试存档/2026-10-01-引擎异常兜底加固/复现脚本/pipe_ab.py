# -*- coding: utf-8 -*-
"""精确对照：管道被提前关闭 vs 读到最后，进程各要多久才退出。

判定标准：只要进程最终能自己退出（不是被杀），就说明没有弹模态框。
沙箱里 PyInstaller onefile 的解包/清理本身很慢，所以要 A/B 对比着看。
"""
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

EXE = sys.argv[1]
HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))   # 仓库根
SRC = os.path.join(REPO, "engine", "test_input_25_new", "input.xlsx")
# 工作区放系统临时目录，别把测试产物写进仓库
BASE = os.path.join(tempfile.gettempdir(), "paike_guard_ab")
os.makedirs(BASE, exist_ok=True)

HL = {
    "schedule": {"periods_per_day": 8, "period_minutes": 40, "morning_start": "08:00",
                 "morning_end": "12:20", "afternoon_start": "14:30", "afternoon_end": "16:55",
                 "long_break_after_period": 3, "long_break_minutes": 30,
                 "eye_break_after_period": 6, "eye_break_minutes": 15,
                 "default_break_minutes": 10},
    "fixed_classes": [{"day": 1, "period": 1, "subject": "班会", "note": ""}],
    "classes": {"total_classes": 25, "arts_start": 1, "arts_end": 4, "science_start": 5,
                "science_end": 25, "gaokao_policy": "新高考3+1+2"},
    "consecutive": {"subjects": ["语文", "数学", "英语"], "math_day": 2, "chinese_day": 3,
                    "english_day": 4, "rule": "", "only_core_subjects": True},
    "solver": {"max_time_seconds": 60, "teacher_candidate_k": 1, "num_plans": 1,
               "min_plans": 1, "teacher_decision": False},
    "extra_constraints": "",
}


def prep(name, drop_chinese=False):
    w = os.path.join(BASE, name)
    if os.path.exists(w):
        shutil.rmtree(w)
    inp, out = os.path.join(w, "input"), os.path.join(w, "output")
    os.makedirs(inp)
    os.makedirs(out)
    shutil.copy(SRC, os.path.join(inp, "input.xlsx"))
    if drop_chinese:
        from openpyxl import load_workbook
        wb = load_workbook(os.path.join(inp, "input.xlsx"))
        ws = wb["教师"]
        heads = [str(c.value).strip() if c.value else "" for c in ws[1]]
        col = heads.index("任教学科") + 1
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=col).value or "").strip() == "语文":
                ws.delete_rows(r, 1)
                break
        wb.save(os.path.join(inp, "input.xlsx"))
        wb.close()
    with open(os.path.join(inp, "requirements.txt"), "w", encoding="utf-8") as f:
        f.write("")
    with open(os.path.join(inp, "hard_limits.json"), "w", encoding="utf-8") as f:
        json.dump(HL, f, ensure_ascii=False, indent=2)
    return inp, out


def run(tag, mode):
    inp, out = prep(tag, drop_chinese=True)
    err = open(os.path.join(BASE, tag + ".err"), "wb")
    t_start = time.time()
    child = subprocess.Popen([EXE, inp, out, tag], stdout=subprocess.PIPE, stderr=err)
    first = child.stdout.readline()               # 等到第一行输出（含 onefile 解包时间）
    t_first = time.time()
    if mode == "close":
        child.stdout.close()                      # 模拟 `| head -1`
    else:
        child.stdout.read()                       # 读到 EOF
    t_after = time.time()
    try:
        rc = child.wait(timeout=300)
        t_exit = time.time()
        killed = False
    except subprocess.TimeoutExpired:
        child.kill()
        rc, t_exit, killed = None, time.time(), True
    err.close()
    print(f"[{tag} / {mode}]")
    print(f"    启动到首行输出 : {t_first - t_start:6.1f}s   （含 onefile 解包）")
    print(f"    首行后到收到EOF: {t_after - t_first:6.1f}s")
    print(f"    到进程真正退出 : {t_exit - t_first:6.1f}s")
    print(f"    退出码={rc}   {'✘ 300 秒仍未退出（疑似弹框）' if killed else '✔ 自己退出了'}")
    st = os.path.join(out, "status.json")
    if os.path.exists(st):
        d = json.load(open(st, encoding="utf-8"))
        print(f"    status.json: code={d['code']}  方案数={len(d['plans'])}")
    else:
        print("    ✘ 没有 status.json")
    err_text = open(os.path.join(BASE, tag + ".err"), "rb").read().decode("utf-8", "replace").strip()
    print(f"    stderr: {err_text[-300:] if err_text else '(空)'}")
    sys.stdout.flush()


run("A_管道提前关闭", "close")
run("B_读到EOF", "drain")
print("对照结束")
