# -*- coding: utf-8 -*-
"""引擎健壮性回归测试：确认异常一律走 status.json，不再弹模态框卡住进程。

用法：python guard_test.py <engine.exe 的路径>

覆盖 5 个场景：
  1. 输入目录不存在        → 期望 status.json code=2，提示能看懂
  2. 目录存在但没有 input.xlsx → 期望 code=2，提示"没有 input.xlsx"
  3. input.xlsx 不是合法工作簿（模拟损坏/被占用） → 期望 code=2
  4. 外壳提前关闭 stdout 管道（就是之前弹框的那个场景） → 期望进程正常退出，不弹框
  5. 真实数据（12 位语文教师，无解） → 期望秒级报"至少需要 13 位"
"""
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

EXE = sys.argv[1]
# 既能测打包好的 exe，也能直接测源码：传 .py 就用 python 跑它
CMD = [r"D:\python\python.exe", EXE] if EXE.lower().endswith(".py") else [EXE]
HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))   # 仓库根（脚本在 测试存档/…/复现脚本/）
SRC_XLSX = os.path.join(REPO, "engine", "test_input_25_new", "input.xlsx")
# 工作区放系统临时目录，别把测试产物写进仓库
BASE = os.path.join(tempfile.gettempdir(), "paike_guard_work")
os.makedirs(BASE, exist_ok=True)

HL = {
    "schedule": {"periods_per_day": 8, "period_minutes": 40, "morning_start": "08:00",
                 "morning_end": "12:20", "afternoon_start": "14:30", "afternoon_end": "16:55",
                 "long_break_after_period": 3, "long_break_minutes": 30,
                 "eye_break_after_period": 6, "eye_break_minutes": 15,
                 "default_break_minutes": 10},
    "fixed_classes": [{"day": 1, "period": 1, "subject": "班会", "note": ""}],
    "classes": {"total_classes": 25, "arts_start": 1, "arts_end": 4, "science_start": 5,
                "science_end": 25, "gaokao_policy": "新高考3+1+2"},
    "consecutive": {"subjects": ["语文", "数学", "英语"], "math_day": 2, "chinese_day": 3,
                    "english_day": 4, "rule": "", "only_core_subjects": True},
    "solver": {"max_time_seconds": 60, "teacher_candidate_k": 1, "num_plans": 1,
               "min_plans": 1, "teacher_decision": False},
    "extra_constraints": "",
}


def make_work(name, kind):
    """kind: 'nodir' | 'empty' | 'brokenxlsx' | 'noplan'  → 返回 (in_dir, out_dir)"""
    w = os.path.join(BASE, name)
    if os.path.exists(w):
        shutil.rmtree(w)
    out = os.path.join(w, "output")
    os.makedirs(out)
    inp = os.path.join(w, "input")
    if kind == "nodir":
        return os.path.join(w, "不存在的目录"), out
    os.makedirs(inp)
    if kind == "brokenxlsx":
        with open(os.path.join(inp, "input.xlsx"), "w", encoding="utf-8") as f:
            f.write("这不是一个 Excel 文件")
    if kind == "noplan":
        shutil.copy(SRC_XLSX, os.path.join(inp, "input.xlsx"))
        from openpyxl import load_workbook
        wb = load_workbook(os.path.join(inp, "input.xlsx"))
        ws = wb["教师"]
        heads = [str(c.value).strip() if c.value else "" for c in ws[1]]
        col = heads.index("任教学科") + 1
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=col).value or "").strip() == "语文":
                ws.delete_rows(r, 1)
                break
        wb.save(os.path.join(inp, "input.xlsx"))
        wb.close()
    with open(os.path.join(inp, "requirements.txt"), "w", encoding="utf-8") as f:
        f.write("")
    with open(os.path.join(inp, "hard_limits.json"), "w", encoding="utf-8") as f:
        json.dump(HL, f, ensure_ascii=False, indent=2)
    return inp, out


def read_status(out):
    p = os.path.join(out, "status.json")
    if not os.path.exists(p):
        return None
    with open(p, encoding="utf-8") as f:
        return json.load(f)


def show(tag, wall, rc, st, extra=""):
    print("=" * 74)
    print(f"[{tag}]  耗时={wall:.1f}s  退出码={rc}")
    if st is None:
        print("  ✘ 没有生成 status.json")
    else:
        print(f"  ✔ code={st['code']}  方案数={len(st['plans'])}")
        print(f"    提示语: {st['message']}")
    if extra:
        print(f"  {extra}")


# ---------- 场景 1~3、5：常规调用 ----------
CASES = [
    ("1_输入目录不存在", "nodir"),
    ("2_目录里没有input.xlsx", "empty"),
    ("3_input.xlsx损坏", "brokenxlsx"),
    ("5_真实数据但无解", "noplan"),
]
for tag, kind in CASES:
    inp, out = make_work(tag, kind)
    t0 = time.time()
    p = subprocess.run(CMD + [inp, out, tag], capture_output=True, timeout=300)
    wall = time.time() - t0
    show(tag, wall, p.returncode, read_status(out))

# ---------- 场景 4：外壳提前关闭 stdout 管道（之前弹模态框的那个） ----------
print("=" * 74)
inp, out = make_work("4_管道提前关闭", "noplan")
t0 = time.time()
child = subprocess.Popen(CMD + [inp, out, "pipe"],
                         stdout=subprocess.PIPE,
                         stderr=open(os.path.join(BASE, "pipe_stderr.txt"), "wb"))
first = child.stdout.readline()          # 只读第一行
child.stdout.close()                     # 然后关掉读端，模拟 `| head -1`
try:
    # 超时给足：沙箱/首次运行时 PyInstaller onefile 解包+清理本身就很慢，
    # 设太短会把"慢"误判成"卡住"。真正的模态框是永远不退出。
    rc = child.wait(timeout=300)
    wall = time.time() - t0
    st = read_status(out)
    show("4_管道提前关闭", wall, rc, st,
         "✔ 进程正常退出（没有弹模态框卡住）" if rc == 0 else "⚠️ 退出码非 0")
except subprocess.TimeoutExpired:
    child.kill()
    print("=" * 74)
    print(f"[4_管道提前关闭]  ✘ 进程 300 秒仍未退出 —— 说明又弹模态框了")
print("=" * 74)
print("全部完成")
