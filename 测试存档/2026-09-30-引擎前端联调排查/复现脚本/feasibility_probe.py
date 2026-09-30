# -*- coding: utf-8 -*-
"""无解判定探针：用仓库自带的合成测试数据，验证「连堂学科的教师数下限」这条规则。

数据来源只有一处：<repo>/engine/test_input_25_new/input.xlsx
（仓库自带、占位命名 T001 / 语文老师001，与任何真实学校数据无关）

做法：把这份数据复制成两个变体——
  A) 原样（语数英各 13 位教师）
  B) 删掉 1 位语文老师（语 12 / 数 13 / 英 13）
然后按外壳的调用契约跑引擎，比较可行性判定结果。

关键配置（写进 hard_limits.json，不改任何代码）：
  teacher_decision=false + teacher_candidate_k=1
  → 问题退化为纯可行性判定，秒级出结果
  （开着候选 3-5 时，同一份无解数据在 60 秒内只能报 UNKNOWN）

预期：
  A → 成功
  B → INFEASIBLE（连堂规则下 25 个班至少需要 13 位语文老师）

运行：python feasibility_probe.py
"""
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from openpyxl import load_workbook

PY = sys.executable
HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
ENGINE_DIR = os.path.join(REPO, "engine")
SRC_XLSX = os.path.join(ENGINE_DIR, "test_input_25_new", "input.xlsx")

# 外壳会写的 hard_limits.json（与 src-tauri/src/main.rs 的 Config::default 对齐）
HARD_LIMITS = {
    "schedule": {
        "periods_per_day": 8, "period_minutes": 40,
        "morning_start": "08:00", "morning_end": "12:20",
        "afternoon_start": "14:30", "afternoon_end": "16:55",
        "long_break_after_period": 3, "long_break_minutes": 30,
        "eye_break_after_period": 6, "eye_break_minutes": 15,
        "default_break_minutes": 10,
    },
    "fixed_classes": [
        {"day": 1, "period": 1, "subject": "班会", "note": "由班主任上课"},
        {"day": 1, "period": 8, "subject": "研究性学习", "note": ""},
        {"day": 4, "period": 8, "subject": "校本课", "note": ""},
    ],
    "classes": {
        "total_classes": 25, "arts_start": 1, "arts_end": 4,
        "science_start": 5, "science_end": 25, "gaokao_policy": "新高考3+1+2",
    },
    "consecutive": {
        "subjects": ["语文", "数学", "英语"],
        "math_day": 2, "chinese_day": 3, "english_day": 4,
        "rule": "", "only_core_subjects": True,
    },
    "solver": {
        "max_time_seconds": 60,
        "teacher_candidate_k": 1,
        "num_plans": 1, "min_plans": 1,
        "teacher_decision": False,
    },
    "extra_constraints": "",
}


def count_teachers_by_subject(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["教师"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    heads = [str(h).strip() if h else "" for h in rows[0]]
    i_sub = heads.index("任教学科")
    cnt = {}
    for r in rows[1:]:
        if not r or all(c is None or str(c).strip() == "" for c in r):
            continue
        s = str(r[i_sub] or "").strip()
        if s:
            cnt[s] = cnt.get(s, 0) + 1
    return cnt


def drop_one_teacher(path, subject):
    """删掉一位指定学科的教师（写入新文件，不动源文件）。"""
    wb = load_workbook(path)
    ws = wb["教师"]
    heads = [str(c.value).strip() if c.value else "" for c in ws[1]]
    i_sub = heads.index("任教学科")
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row=row, column=i_sub + 1).value or "").strip() == subject:
            ws.delete_rows(row, 1)
            break
    wb.save(path)
    wb.close()


def run_variant(name, xlsx, base):
    work = os.path.join(base, name)
    os.makedirs(os.path.join(work, "input"))
    os.makedirs(os.path.join(work, "output"))
    shutil.copy(xlsx, os.path.join(work, "input", "input.xlsx"))
    with open(os.path.join(work, "input", "requirements.txt"), "w", encoding="utf-8") as f:
        f.write("")
    with open(os.path.join(work, "input", "hard_limits.json"), "w", encoding="utf-8") as f:
        json.dump(HARD_LIMITS, f, ensure_ascii=False, indent=2)

    t0 = time.time()
    subprocess.run(
        [PY, os.path.join(ENGINE_DIR, "engine.py"),
         os.path.join(work, "input"), os.path.join(work, "output"), name],
        cwd=ENGINE_DIR, capture_output=True, text=True,
        encoding="utf-8", errors="replace",
        env=dict(os.environ, PYTHONIOENCODING="utf-8"),
    )
    wall = time.time() - t0
    st = os.path.join(work, "output", "status.json")
    code, nplan, msg = "无 status.json", 0, ""
    if os.path.exists(st):
        d = json.load(open(st, encoding="utf-8"))
        code, nplan, msg = d["code"], len(d["plans"]), d.get("message", "")
    return wall, code, nplan, msg


def main():
    if not os.path.exists(SRC_XLSX):
        print("找不到仓库测试数据:", SRC_XLSX)
        return 1

    print("数据来源:", SRC_XLSX)
    print()

    with tempfile.TemporaryDirectory(prefix="paike_probe_") as base:
        # 变体 A：原样
        a = os.path.join(base, "A_原样.xlsx")
        shutil.copy(SRC_XLSX, a)

        # 变体 B：删掉 1 位语文老师
        b = os.path.join(base, "B_少一位语文.xlsx")
        shutil.copy(SRC_XLSX, b)
        drop_one_teacher(b, "语文")

        print("两个变体的语数英教师数：")
        for tag, p in (("A 原样", a), ("B 少一位语文", b)):
            c = count_teachers_by_subject(p)
            print(f"  {tag:12s} 语文={c.get('语文', 0)}  数学={c.get('数学', 0)}  "
                  f"英语={c.get('英语', 0)}   连堂判据需 ≥13 → "
                  f"{'满足' if c.get('语文', 0) >= 13 else '不满足'}")
        print()

        print("跑引擎（teacher_decision=false, 候选=1）：")
        results = []
        for tag, p in (("A_原样", a), ("B_少一位语文", b)):
            wall, code, nplan, msg = run_variant(tag, p, base)
            verdict = {0: "✅ 成功", 1: "✅ 成功（有警告）", 2: "❌ 失败"}.get(code, str(code))
            print(f"  {tag:14s} 耗时={wall:5.1f}s   code={code}  "
                  f"方案数={nplan}   {verdict}")
            print(f"      提示语: {msg}")
            results.append((tag, wall, code))

    print()
    a_ok = results[0][2] in (0, 1)
    b_bad = results[1][2] == 2
    if a_ok and b_bad:
        print("结论成立：只差 1 位语文老师，结果从「可行」变成「无解」。")
        print("→ 连堂规则下，连堂学科的教师数必须 ≥ ceil(班数 / 2) = 13（25 个班）。")
        return 0
    print("结论与预期不符，需要重新检查。")
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
